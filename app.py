from flask import Flask, render_template, request, redirect, url_for, send_file, flash
import sqlite3
import datetime
import io
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = "skinet_secret_key"

DATA_DIR = "/data" if os.path.exists("/data") else "."
DB_NAME = os.path.join(DATA_DIR, "database.db")
# DB_NAME = "database.db"

# Listas maestras para formularios
SEDES = [
    "Sede 01 - Piura", "Sede 02 - Trujillo", "Sede 03 - Lima", "Sede 04 - Arequipa 1",
    "Sede 05 - Arequipa 2", "Sede 06 - Cuzco Infancia", "Sede 07 - Pucallpa", 
    "Sede 08 - Cajamarca", "Sede 09 - Tingo Maria", "Sede 10 - Iquitos", "Sede 11 - Andahuaylas",
    "Sede 12 - Huancayo", "Sede 13 - Puerto Maldonado", "Sede 14 - Nazca", "Sede 15 - Ica", "Sede 16 - Tacna"
]
SISTEMAS = [
    "Sistema PV - Off Grid", "Sistema PV - On Grid", "Sistema PV - Híbrido",
    "Bombeo Solar", "Terma Solar"
]
ACTIVIDADES = ["Visita Técnica", "Instalación", "Revisión Técnica"]
ESTADOS = ["Finalizado Conforme", "Pendiente Visto Bueno", "Incompleto / Faltan Equipos", "Reprogramado"]

# ----------------------------------------------------
# Inicialización de la Base de Datos
# ----------------------------------------------------
def init_db():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS registro_campo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            sede TEXT NOT NULL,
            tecnico TEXT NOT NULL,
            cliente TEXT NOT NULL,
            tipo_actividad TEXT NOT NULL,
            sistema TEXT NOT NULL,
            equipos TEXT NOT NULL,
            hora_inicio TEXT NOT NULL,
            hora_fin TEXT NOT NULL,
            tiempo_horas REAL,
            llamada_vb TEXT NOT NULL,
            ticket_ot TEXT,
            estado_final TEXT NOT NULL,
            observaciones TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

# ----------------------------------------------------
# Rutas Principales
# ----------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html", sedes=SEDES, sistemas=SISTEMAS, actividades=ACTIVIDADES, estados=ESTADOS)

@app.route("/guardar", methods=["POST"])
def guardar():
    fecha = request.form.get("fecha") or datetime.date.today().strftime("%Y-%m-%d")
    sede = request.form.get("sede")
    tecnico = request.form.get("tecnico")
    cliente = request.form.get("cliente")
    tipo_actividad = request.form.get("tipo_actividad")
    sistema = request.form.get("sistema")
    equipos = request.form.get("equipos")
    hora_inicio = request.form.get("hora_inicio")
    hora_fin = request.form.get("hora_fin")
    llamada_vb = request.form.get("llamada_vb")
    ticket_ot = request.form.get("ticket_ot")
    estado_final = request.form.get("estado_final")
    observaciones = request.form.get("observaciones")

    tiempo_horas = 0.0
    try:
        t1 = datetime.datetime.strptime(hora_inicio, "%H:%M")
        t2 = datetime.datetime.strptime(hora_fin, "%H:%M")
        tiempo_horas = round((t2 - t1).seconds / 3600.0, 2)
    except Exception:
        tiempo_horas = 0.0

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO registro_campo 
        (fecha, sede, tecnico, cliente, tipo_actividad, sistema, equipos, hora_inicio, hora_fin, tiempo_horas, llamada_vb, ticket_ot, estado_final, observaciones)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', (fecha, sede, tecnico, cliente, tipo_actividad, sistema, equipos, hora_inicio, hora_fin, tiempo_horas, llamada_vb, ticket_ot, estado_final, observaciones))
    conn.commit()
    conn.close()

    flash("¡Registro guardado exitosamente!", "success")
    return redirect(url_for("dashboard", fecha=fecha))

@app.route("/dashboard")
def dashboard():
    fecha_filtro = request.args.get("fecha") or datetime.date.today().strftime("%Y-%m-%d")
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM registro_campo WHERE fecha = ? ORDER BY id DESC", (fecha_filtro,))
    registros = cursor.fetchall()
    conn.close()

    return render_template("dashboard.html", registros=registros, fecha_filtro=fecha_filtro)

# ----------------------------------------------------
# Editar y Eliminar Registros (NUEVO)
# ----------------------------------------------------
@app.route("/editar/<int:id>", methods=["GET", "POST"])
def editar(id):
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    if request.method == "POST":
        fecha = request.form.get("fecha")
        sede = request.form.get("sede")
        tecnico = request.form.get("tecnico")
        cliente = request.form.get("cliente")
        tipo_actividad = request.form.get("tipo_actividad")
        sistema = request.form.get("sistema")
        equipos = request.form.get("equipos")
        hora_inicio = request.form.get("hora_inicio")
        hora_fin = request.form.get("hora_fin")
        llamada_vb = request.form.get("llamada_vb")
        ticket_ot = request.form.get("ticket_ot")
        estado_final = request.form.get("estado_final")
        observaciones = request.form.get("observaciones")

        tiempo_horas = 0.0
        try:
            t1 = datetime.datetime.strptime(hora_inicio, "%H:%M")
            t2 = datetime.datetime.strptime(hora_fin, "%H:%M")
            tiempo_horas = round((t2 - t1).seconds / 3600.0, 2)
        except Exception:
            tiempo_horas = 0.0

        cursor.execute('''
            UPDATE registro_campo 
            SET fecha=?, sede=?, tecnico=?, cliente=?, tipo_actividad=?, sistema=?, equipos=?, 
                hora_inicio=?, hora_fin=?, tiempo_horas=?, llamada_vb=?, ticket_ot=?, estado_final=?, observaciones=?
            WHERE id=?
        ''', (fecha, sede, tecnico, cliente, tipo_actividad, sistema, equipos, hora_inicio, hora_fin, tiempo_horas, llamada_vb, ticket_ot, estado_final, observaciones, id))
        conn.commit()
        conn.close()

        flash("¡Registro actualizado correctamente!", "info")
        return redirect(url_for("dashboard", fecha=fecha))

    # Cargar registro actual para la vista de edición
    cursor.execute("SELECT * FROM registro_campo WHERE id = ?", (id,))
    registro = cursor.fetchone()
    conn.close()

    if not registro:
        flash("El registro no existe.", "danger")
        return redirect(url_for("dashboard"))

    return render_template("editar.html", r=registro, sedes=SEDES, sistemas=SISTEMAS, actividades=ACTIVIDADES, estados=ESTADOS)

@app.route("/eliminar/<int:id>", methods=["POST"])
def eliminar(id):
    fecha_filtro = request.form.get("fecha_filtro") or datetime.date.today().strftime("%Y-%m-%d")
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("DELETE FROM registro_campo WHERE id = ?", (id,))
    conn.commit()
    conn.close()

    flash("El registro fue eliminado.", "warning")
    return redirect(url_for("dashboard", fecha=fecha_filtro))

# ----------------------------------------------------
# Exportación a Excel en Memoria (SOLUCIONA FILE NOT FOUND)
# ----------------------------------------------------
@app.route("/exportar_excel")
def exportar_excel():
    fecha_filtro = request.args.get("fecha") or datetime.date.today().strftime("%Y-%m-%d")
    
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM registro_campo WHERE fecha = ?", (fecha_filtro,))
    data = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Reporte_{fecha_filtro}"
    ws.views.sheetView[0].showGridLines = True

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    headers = [
        "ID", "Fecha", "Sede", "Técnico", "Cliente / Proyecto", 
        "Tipo Actividad", "Sistema", "Equipos / Materiales", 
        "Hora Inicio", "Hora Fin", "Horas Inst.", "Llamada VB", 
        "N° OT/Ticket", "Estado Final", "Observaciones"
    ]
    
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in data:
        ws.append(list(row))

    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), min_col=1, max_col=15):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 2, 9, 10, 11, 12, 13]:
                cell.alignment = Alignment(horizontal="center")

    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Enviar archivo directamente desde memoria RAM (evita FileNotFoundError)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        download_name=f"Reporte_Campo_{fecha_filtro}.xlsx",
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)